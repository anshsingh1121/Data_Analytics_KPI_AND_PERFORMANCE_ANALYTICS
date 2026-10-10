import logging
import os
import pandas as pd
from typing import Dict, Any, Optional
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

class ExcelExporter:
    """
    Generates a consolidated, professionally styled Excel report workbook
    containing all management summaries, KPI calculations, cycle comparisons,
    and operational drilldowns.
    """

    def __init__(self, output_dir: str):
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)

    def export(
        self,
        sheets_data: Dict[str, pd.DataFrame],
        metadata: Optional[Dict[str, Any]] = None,
        filename_prefix: str = "Data_Certification_KPI_Report"
    ) -> str:
        """
        Exports multiple DataFrames into a single, multi-tab Excel workbook with professional styling.
        """
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"{filename_prefix}_{timestamp}.xlsx"
        filepath = os.path.join(self.output_dir, filename)

        # Also write a canonical copy for immediate download from the dashboard
        canonical_path = os.path.join(self.output_dir, f"{filename_prefix}_latest.xlsx")

        try:
            wb = openpyxl.Workbook()
            # Remove default first sheet
            wb.remove(wb.active)

            # Header and styling constants
            header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            title_font = Font(name="Calibri", size=14, bold=True, color="1F497D")
            meta_font = Font(name="Calibri", size=9, italic=True, color="595959")
            border_thin = Border(
                left=Side(style='thin', color='D9D9D9'),
                right=Side(style='thin', color='D9D9D9'),
                top=Side(style='thin', color='D9D9D9'),
                bottom=Side(style='thin', color='D9D9D9')
            )

            for sheet_name, df in sheets_data.items():
                if df is None:
                    continue

                # Ensure sheet name is max 31 characters
                clean_sheet_name = str(sheet_name)[:31]
                ws = wb.create_sheet(title=clean_sheet_name)
                ws.views.sheetView[0].showGridLines = True

                # Write sheet title and metadata
                ws.cell(row=1, column=1, value=f"Data Certification Performance — {clean_sheet_name}").font = title_font
                
                meta_str = f"Exported: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Mode: {metadata.get('mode', 'Demo') if metadata else 'Demo'}"
                if metadata and 'cycle_name' in metadata:
                    meta_str += f" | Cycle: {metadata['cycle_name']}"
                ws.cell(row=2, column=1, value=meta_str).font = meta_font

                start_row = 4

                if df.empty:
                    ws.cell(row=start_row, column=1, value="No records available for this section.")
                    continue

                # Write header row
                for col_idx, col_name in enumerate(df.columns, 1):
                    cell = ws.cell(row=start_row, column=col_idx, value=str(col_name))
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.border = border_thin

                # Write data rows
                for r_idx, row in enumerate(df.itertuples(index=False), start_row + 1):
                    for c_idx, val in enumerate(row, 1):
                        cell = ws.cell(row=r_idx, column=c_idx, value=val)
                        cell.border = border_thin
                        # Formatting numbers
                        if isinstance(val, (int, float)):
                            cell.alignment = Alignment(horizontal="right")
                        else:
                            cell.alignment = Alignment(horizontal="left")

                # Auto-fit column widths
                for col in ws.columns:
                    col_letter = get_column_letter(col[0].column)
                    max_len = max(len(str(cell.value or '')) for cell in col)
                    ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

                # Add autofilter
                last_col_letter = get_column_letter(len(df.columns))
                ws.auto_filter.ref = f"A{start_row}:{last_col_letter}{start_row + len(df)}"

            wb.save(filepath)
            wb.save(canonical_path)
            logger.info(f"Excel report successfully generated at {filepath}")
            return filepath

        except Exception as e:
            logger.error(f"Failed to export Excel workbook: {e}")
            return ""
